"""E2E: /statistics/ hub + diagnostics delete/export on client card."""
from __future__ import annotations

import io
import re
from datetime import date, datetime, time, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.auth.passwords import hash_password
from app.database import Base, get_async_db
from app.diagnostics.catalog import BHS
from app.models import Booking, Calendar, Category, ClientCard, Consultant, DiagnosticAttempt, Service, User
from app.services.diagnostics_service import complete_attempt, start_attempt


@pytest.fixture()
def stats_client(tmp_path):
    from app.main import app

    engine = create_async_engine(
        "sqlite+aiosqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async def _prepare():
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
            try:
                await conn.execute(text("ALTER TABLE consultants ADD COLUMN public_slug VARCHAR(64)"))
            except Exception:
                pass

        async with session_factory() as db:
            cat = Category(name_category="Общая")
            db.add(cat)
            await db.flush()
            spec_user = User(
                username="stats-spec@test.com",
                email="stats-spec@test.com",
                password=hash_password("specpass"),
                is_active=True,
                date_joined=datetime.utcnow(),
            )
            client_user = User(
                username="+79991112233",
                email="client-stats@test.com",
                password=hash_password("clientpass"),
                is_active=True,
                date_joined=datetime.utcnow(),
            )
            db.add(spec_user)
            db.add(client_user)
            await db.flush()
            consultant = Consultant(
                first_name="Стат",
                last_name="Спец",
                email="stats-spec@test.com",
                phone="+79990001122",
                category_of_specialist_id=cat.id,
                user_id=spec_user.id,
            )
            db.add(consultant)
            await db.flush()
            await db.execute(
                text("UPDATE consultants SET public_slug = :slug WHERE id = :id"),
                {"slug": "statspec", "id": consultant.id},
            )
            cal = Calendar(
                consultant_id=consultant.id,
                name="Основной",
                color="#111111",
                is_active=True,
            )
            db.add(cal)
            await db.flush()
            service = Service(
                consultant_id=consultant.id,
                calendar_id=cal.id,
                name="Консультация",
                duration_minutes=60,
                is_active=True,
            )
            db.add(service)
            await db.flush()
            card = ClientCard(
                consultant_id=consultant.id,
                client_user_id=client_user.id,
                name="Иван Клиентов",
                phone="+79991112233",
                email="client-stats@test.com",
                telegram="ivan_client",
                notes="Тестовая заметка",
            )
            db.add(card)
            await db.flush()
            today = date.today()
            for status, day_offset in (
                ("completed", -2),
                ("cancelled", -1),
                ("confirmed", 0),
                ("pending", 1),
            ):
                db.add(
                    Booking(
                        calendar_id=cal.id,
                        service_id=service.id,
                        client_card_id=card.id,
                        client_name="Иван Клиентов",
                        client_phone="+79991112233",
                        client_email="client-stats@test.com",
                        client_telegram="ivan_client",
                        booking_date=today + timedelta(days=day_offset),
                        booking_time=time(12, 0),
                        status=status,
                        source="specialist",
                    )
                )
            answers = {item.id: "1" for item in BHS.items}
            attempt = await start_attempt(
                db,
                client_user_id=client_user.id,
                consultant_id=consultant.id,
                test_code=BHS.code,
                source="cabinet",
                client_card_id=card.id,
            )
            await complete_attempt(db, attempt=attempt, answers=answers)
            await db.commit()
            return consultant.id, card.id, attempt.id, client_user.id

    consultant_id, card_id, attempt_id, client_user_id = __import__("asyncio").run(_prepare())

    async def override_get_async_db():
        async with session_factory() as db:
            yield db

    app.dependency_overrides[get_async_db] = override_get_async_db

    import app.database as database_module

    database_module._async_engine = engine
    database_module._AsyncSessionLocal = session_factory

    client = TestClient(app)
    try:
        yield client, consultant_id, card_id, attempt_id, client_user_id, session_factory
    finally:
        app.dependency_overrides.clear()
        database_module._async_engine = None
        database_module._AsyncSessionLocal = None
        __import__("asyncio").run(engine.dispose())


def _login_spec(client: TestClient) -> None:
    page = client.get("/login/", follow_redirects=True)
    m = re.search(r'name="csrf_token"\s+value="([^"]+)"', page.text)
    assert m
    client.post(
        "/login/",
        data={"login": "stats-spec@test.com", "password": "specpass", "csrf_token": m.group(1)},
        follow_redirects=True,
    )


def test_statistics_page_and_xlsx(stats_client):
    client, *_ = stats_client
    _login_spec(client)

    booking_page = client.get("/booking/", follow_redirects=True)
    assert booking_page.status_code == 200
    assert 'id="bookings-dashboard"' not in booking_page.text

    clients_page = client.get("/clients/", follow_redirects=True)
    assert clients_page.status_code == 200
    assert 'id="clients-dashboard"' not in clients_page.text

    nav = clients_page.text + booking_page.text
    assert "/statistics/" in nav
    assert "Статистика" in nav

    today = date.today()
    start = today.replace(day=1)
    stats = client.get(
        f"/statistics/?from={start.isoformat()}&to={(today + timedelta(days=7)).isoformat()}",
        follow_redirects=True,
    )
    assert stats.status_code == 200, stats.text[:400]
    assert 'id="stats-dashboard"' in stats.text
    assert "Завершены" in stats.text
    assert "Отменены" in stats.text
    assert "Выгрузить Excel" in stats.text

    xlsx = client.get(
        f"/statistics/export.xlsx?from={start.isoformat()}&to={(today + timedelta(days=7)).isoformat()}"
    )
    assert xlsx.status_code == 200
    assert "spreadsheetml" in xlsx.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(xlsx.content))
    ws = wb.active
    assert ws["A1"].value == "Статистика консультаций"
    headers = [ws.cell(row=5, column=c).value for c in range(1, 7)]
    assert "Дата" in headers and "Клиент" in headers


def test_diagnostics_delete_and_export_xlsx(stats_client):
    client, _cid, card_id, attempt_id, _uid, session_factory = stats_client
    _login_spec(client)

    detail = client.get(f"/clients/{card_id}/", follow_redirects=True)
    assert detail.status_code == 200
    assert "Выгрузить Excel" in detail.text
    assert "Удалить" in detail.text
    assert f"/clients/{card_id}/diagnostics/{attempt_id}/delete/" in detail.text

    export = client.get(f"/clients/{card_id}/diagnostics/export.xlsx")
    assert export.status_code == 200
    assert "spreadsheetml" in export.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(export.content))
    assert "Профиль" in wb.sheetnames
    assert "Диагностика — сводка" in wb.sheetnames
    assert "Шкалы" in wb.sheetnames
    assert "Анамнез" in wb.sheetnames
    profile = wb["Профиль"]
    assert profile["B4"].value == "Иван Клиентов"
    summary = wb["Диагностика — сводка"]
    assert summary.cell(row=4, column=3).value == BHS.code

    page = client.get(f"/clients/{card_id}/", follow_redirects=True)
    csrf = re.search(r'name="csrf_token"\s+value="([^"]+)"', page.text)
    assert csrf
    deleted = client.post(
        f"/clients/{card_id}/diagnostics/{attempt_id}/delete/",
        data={"csrf_token": csrf.group(1)},
        follow_redirects=True,
    )
    assert deleted.status_code == 200
    assert f'data-attempt-id="{attempt_id}"' not in deleted.text

    import asyncio

    async def _count():
        async with session_factory() as db:
            row = await db.get(DiagnosticAttempt, attempt_id)
            return row

    assert asyncio.run(_count()) is None
